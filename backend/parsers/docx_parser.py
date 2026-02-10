"""
Word document parser for extracting events from .docx files.

Uses python-docx to read Word documents and extract text from paragraphs
and tables, then uses AI via BaseParser to extract event information.
"""

from docx import Document
import os
from parsers.base_parser import BaseParser


class DocxParser(BaseParser):
    """
    Parser for Microsoft Word documents (.docx).
    
    Extracts text from paragraphs and tables, preserving structure,
    then uses AI to identify events, dates, and module information.
    """
    
    def parse(self, file_path: str) -> list:
        """
        Parse a Word document and extract events.
        
        Args:
            file_path: Path to the Word document (.docx)
            
        Returns:
            List of event dictionaries with module, title, date, type, description
            
        Raises:
            FileNotFoundError: If the file doesn't exist
            Exception: If parsing fails
        """
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"Word document not found: {file_path}")
        
        try:
            doc = Document(file_path)
            text_parts = []
            
            # Extract paragraphs
            for paragraph in doc.paragraphs:
                text = paragraph.text.strip()
                if text:
                    text_parts.append(text)
            
            # Extract tables
            for table in doc.tables:
                text_parts.append("\\n--- Table ---")
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells)
                    if row_text.strip():
                        text_parts.append(row_text)
                text_parts.append("--- End Table ---\\n")
            
            combined_text = "\\n".join(text_parts)
            
            if not combined_text.strip():
                print("WARNING: Word document appears to be empty")
                return []
            
            # Use AI extraction from base parser
            events = self.extract_events_with_ai(combined_text)
            
            return events
            
        except Exception as e:
            print(f"ERROR: Failed to parse Word document: {e}")
            if hasattr(self, 'config') and self.config.DEBUG:
                import traceback
                traceback.print_exc()
            raise
